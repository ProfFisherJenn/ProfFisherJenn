import os
import sys
from pathlib import Path
from datetime import datetime
from shutil import copyfile
from docx import Document
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

# ==============================
# Utilities
# ==============================
def log(msg: str):
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    line = f"[{stamp}] {msg}"
    print(line)
    with open(LOG_PATH, "a", encoding="utf-8") as f:
        f.write(line + "\n")

def pause():
    try:
        if os.name == "nt":
            os.system("pause")
        else:
            input("Press Enter to exit...")
    except Exception:
        pass

def open_folder(path: Path):
    try:
        if os.name == "nt":
            os.startfile(str(path))
        elif sys.platform == "darwin":
            os.system(f'open "{path}"')
        else:
            os.system(f'xdg-open "{path}"')
    except Exception as e:
        log(f"Could not open folder {path}: {e}")

# ==============================
# Paths
# ==============================
BASE_DIR = Path(__file__).resolve().parent
LOG_PATH = BASE_DIR / "conversion_log.txt"
INPUT_DIR = BASE_DIR / "input_docs"
OUTPUT_DIR = BASE_DIR / "output_pptx"
INPUT_DIR.mkdir(exist_ok=True)
OUTPUT_DIR.mkdir(exist_ok=True)

try:
    LOG_PATH.unlink(missing_ok=True)
except Exception:
    pass

log("=== Word → PowerPoint Converter (v2) started ===")
log(f"Script folder: {BASE_DIR}")

# ==============================
# Template discovery (.pptx or .potx)
# ==============================
def find_template(folder: Path) -> Path:
    """Prefer a .pptx. If none, take the newest .potx and clone it to a .pptx for use."""
    pptx_files = sorted(folder.glob("*.pptx"), key=lambda p: p.stat().st_mtime, reverse=True)
    if pptx_files:
        return pptx_files[0]

    potx_files = sorted(folder.glob("*.potx"), key=lambda p: p.stat().st_mtime, reverse=True)
    if potx_files:
        potx = potx_files[0]
        working = folder / "_working_template_from_potx.pptx"
        try:
            copyfile(potx, working)
            log(f"Cloned template: {potx.name} -> {working.name}")
            return working
        except Exception as e:
            raise FileNotFoundError(f"Found .potx but could not clone to .pptx: {e}")

    raise FileNotFoundError(
        f"No template found. Place a .pptx or .potx next to this script in {folder}."
    )

# ==============================
# Helpers
# ==============================
def find_layout(prs: Presentation, target_name: str):
    name_l = target_name.lower()
    # exact match
    for layout in prs.slide_layouts:
        if layout.name and layout.name.lower() == name_l:
            return layout
    # substring match
    for layout in prs.slide_layouts:
        if layout.name and name_l in layout.name.lower():
            return layout
    raise ValueError(f"Slide layout '{target_name}' not found in template.")

def find_placeholder_by_type(slide, ph_type: PP_PLACEHOLDER):
    for ph in slide.placeholders:
        try:
            if ph.placeholder_format.type == ph_type:
                return ph
        except Exception:
            continue
    return None

def find_placeholder_fallback(slide, keyword: str):
    k = keyword.lower()
    for ph in slide.placeholders:
        name = getattr(ph, "name", "") or ""
        if k in name.lower():
            return ph
    return None

def fill_textframe(text_frame, paras, map_level_fn=None):
    text_frame.clear()
    first_para = text_frame.paragraphs[0]
    used_first = False

    for para in paras:
        if not para.text.strip():
            continue

        p = first_para if not used_first else text_frame.add_paragraph()
        used_first = True

        level, bullet = (0, False)
        if map_level_fn is not None:
            level, bullet = map_level_fn(para)

        p.level = level
        p.bullet = bullet

        for run in para.runs:
            new_run = p.add_run()
            new_run.text = run.text
            f = new_run.font
            if run.bold:
                f.bold = True
            if run.italic:
                f.italic = True
            if run.underline:
                f.underline = True

def body_level_map(para):
    style = para.style.name
    if style == "List Paragraph":
        return (1, True)   # "Second level" in Slide Master (first bullet level)
    return (0, False)

# ==============================
# Slide assembly
# ==============================
def add_slide(prs: Presentation, layout_name: str, title_text: str,
              subtitle_paras, body_paras):
    layout = find_layout(prs, layout_name)
    slide = prs.slides.add_slide(layout)

    if slide.shapes.title:
        slide.shapes.title.text = title_text or ""

    subtitle_ph = (find_placeholder_by_type(slide, PP_PLACEHOLDER.SUBTITLE)
                   or find_placeholder_fallback(slide, "Subtitle"))
    if subtitle_ph and subtitle_paras:
        fill_textframe(subtitle_ph.text_frame, subtitle_paras)

    content_ph = (find_placeholder_by_type(slide, PP_PLACEHOLDER.BODY)
                  or find_placeholder_fallback(slide, "Content"))
    if content_ph and body_paras:
        fill_textframe(content_ph.text_frame, body_paras, body_level_map)

def convert_one(doc_path: Path, template_path: Path, out_path: Path):
    log(f"Converting: {doc_path.name}")
    prs = Presentation(template_path)
    doc = Document(doc_path)

    paragraphs = doc.paragraphs
    current_layout = None
    current_title = None
    current_subtitle = []
    current_body = []

    def flush():
        nonlocal current_layout, current_title, current_subtitle, current_body
        if current_layout and current_title is not None:
            add_slide(prs, current_layout, current_title, current_subtitle, current_body)
        current_layout = None
        current_title = None
        current_subtitle = []
        current_body = []

    for para in paragraphs:
        style = para.style.name

        if style == "Title":
            flush()
            current_layout = "Title Slide"
            current_title = para.text.strip()

        elif style.startswith("Heading 1"):
            flush()
            current_layout = "Section Header"
            current_title = para.text.strip()

        elif style.startswith("Heading 2"):
            flush()
            current_layout = "Title and Content"
            current_title = para.text.strip()

        else:
            if style == "Subtitle":
                current_subtitle.append(para)
            else:
                current_body.append(para)

    flush()
    prs.save(out_path)
    log(f"Saved: {out_path.name}")

# ==============================
# Main
# ==============================
def main():
    try:
        template_path = find_template(BASE_DIR)
        log(f"Using template: {template_path.name}")
    except FileNotFoundError as e:
        log(str(e))
        open_folder(BASE_DIR)
        pause()
        return

    docs = sorted(INPUT_DIR.glob("*.docx"))
    if not docs:
        log(f"No .docx files found in: {INPUT_DIR}")
        open_folder(INPUT_DIR)
        pause()
        return

    for doc_path in docs:
        out_path = OUTPUT_DIR / (doc_path.stem + ".pptx")
        try:
            convert_one(doc_path, template_path, out_path)
            os.remove(doc_path)
            log(f"Removed source: {doc_path.name}")
        except Exception as ex:
            log(f"Failed: {doc_path.name} — {ex}")

    log("All done. Opening output folder …")
    open_folder(OUTPUT_DIR)
    pause()

if __name__ == "__main__":
    main()
